import copy
import io
import time
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select
from openpyxl import load_workbook
from backend.main import create_app
from backend.storage import state,versions,audit,uid,now
from backend.solver import solve
from backend.exports import export
from .test_solver import small

@pytest.fixture
def client(tmp_path):
    app=create_app('sqlite:///'+str(tmp_path/'test.db'))
    with TestClient(app) as c:
        c.post('/api/auth/setup',json={'password':'Initial-password-42'})
        token=c.post('/api/auth/login',json={'password':'Initial-password-42'}).json()['csrf']
        c.headers['x-csrf-token']=token
        assert c.post('/api/auth/password',json={'current':'Initial-password-42','password':'Permanent-password-43'}).status_code==200
        token=c.post('/api/auth/login',json={'password':'Permanent-password-43'}).json()['csrf'];c.headers['x-csrf-token']=token
        yield c

def seed(client):
    d=small();r=solve(d,seconds=1);assert r['status']=='complete';vid=uid()
    with client.app.state.engine.begin() as c:
        c.execute(state.update().values(payload=d,revision=1))
        c.execute(versions.insert().values(id=vid,created=now(),status='draft',payload={'snapshot':d,'revision':1,'week':'','lessons':r['lessons'],'quality':r['quality'],'complete':True}))
    return vid,r,d

def test_unauthorized_and_forced_password_change(tmp_path):
    app=create_app('sqlite:///'+str(tmp_path/'auth.db'))
    with TestClient(app) as c:
        assert c.get('/api/data').status_code==401
        assert c.post('/api/generate',json={}).status_code==401
        assert c.get('/api/audit').status_code==401
        assert c.post('/api/auth/setup',json={'password':'Initial-password-42'}).status_code==201
        result=c.post('/api/auth/login',json={'password':'Initial-password-42'}); assert result.status_code==200
        assert 'HttpOnly' in result.headers['set-cookie'];assert 'SameSite=strict' in result.headers['set-cookie']
        assert c.get('/api/data').status_code==403

def test_csrf_validation_and_concurrent_write(client):
    revision=client.get('/api/data').json()['revision'];d=small();record=d['subjects'][0]
    response=client.post('/api/data/subjects',json={'record':record,'revision':revision});assert response.status_code==200
    assert client.post('/api/data/subjects',json={'record':record,'revision':revision}).status_code==409
    assert client.post('/api/data/subjects',headers={'x-csrf-token':'bad'},json={'record':record,'revision':1}).status_code==403
    assert client.post('/api/data/subjects',headers={'Origin':'https://evil.example'},json={'record':record,'revision':1}).status_code==403
    assert client.post('/api/data/subjects',json={'record':{**record,'admin':True},'revision':1}).status_code==422

def test_invalid_manual_move_never_saves(client):
    vid,r,d=seed(client);bad=copy.deepcopy(r['lessons']);bad[1]['slot']=bad[0]['slot']
    assert client.post(f'/api/versions/{vid}/edit',json={'lessons':bad}).status_code==422
    assert len(client.get('/api/versions').json())==1

def test_lifecycle_snapshot_restore_compare(client):
    vid,r,d=seed(client)
    for status in ['review','approved','published']: assert client.post(f'/api/versions/{vid}/status',json={'status':status}).status_code==200
    lessons=copy.deepcopy(r['lessons']);lessons[0]['lock']='locked'
    new=client.post(f'/api/versions/{vid}/edit',json={'lessons':lessons});assert new.status_code==201
    assert client.get(f'/api/versions/{vid}').json()['lessons']==r['lessons']
    diff=client.get('/api/compare',params={'before':vid,'after':new.json()['id']}).json();assert len(diff['changed'])==1
    restored=client.post(f'/api/versions/{vid}/restore');assert restored.status_code==201
    assert client.get('/api/versions/'+restored.json()['id']).json()['status']=='draft'

def test_stale_data_blocks_publication(client):
    vid,r,d=seed(client)
    client.post('/api/data/subjects',json={'revision':1,'record':{**d['subjects'][0],'name':'Новое название'}})
    assert client.post(f'/api/versions/{vid}/status',json={'status':'review'}).status_code==409

def test_exports_use_snapshot_filter_and_escape_formulas(client):
    vid,r,d=seed(client)
    response=client.get(f'/api/versions/{vid}/export/csv?class_id=c0');assert response.status_code==200
    assert '5А' in response.text and '5Б' not in response.text
    assert client.get(f'/api/versions/{vid}/export/pdf').content.startswith(b'%PDF')
    xlsx=client.get(f'/api/versions/{vid}/export/xlsx');wb=load_workbook(io.BytesIO(xlsx.content));assert set(wb.sheetnames)=={'Общая сетка','Классы','Учителя','Кабинеты','Конфликты'}
    d['subjects'][0]['name']='=HYPERLINK("evil")'
    content,_=export({'snapshot':d,'week':'','lessons':r['lessons']},vid,'draft','csv',{})
    assert "'=HYPERLINK" in content.decode('utf-8-sig')

def test_audit_does_not_leak_passwords(client):
    seed(client);rows=client.get('/api/audit').json();assert rows
    assert 'password' not in str(rows).lower();assert 'Permanent' not in str(rows)
    assert client.delete('/api/audit').status_code==405

def test_rate_limit(tmp_path):
    with TestClient(create_app('sqlite:///'+str(tmp_path/'rate.db'))) as c:
        for _ in range(5):assert c.post('/api/auth/login',json={'password':'bad'}).status_code==401
        assert c.post('/api/auth/login',json={'password':'bad'}).status_code==429

def test_generation_accept_and_repair_preserves_untouched(client):
    vid,r,d=seed(client)
    # Add an alternative, then make just the first day's teacher unavailable.
    for p in d['plans']:p['teacher_ids']=['t0','t1']
    with client.app.state.engine.begin() as c:c.execute(state.update().values(payload=d))
    day=r['lessons'][0]['slot']//8
    response=client.post('/api/generate',json={'source_id':vid,'repair_scope':'teachers','repair_resource':'t0','repair_days':[day],'expand_percent':0,'seconds':1})
    assert response.status_code==202,response.text
    jid=response.json()['id']
    for _ in range(100):
        job=client.get('/api/jobs/'+jid).json()
        if job['status']!='running':break
        time.sleep(.05)
    assert job['status']=='complete',job
    accepted=client.post(f'/api/jobs/{jid}/accept',json={});assert accepted.status_code==201,accepted.text
    new=client.get('/api/versions/'+accepted.json()['id']).json()
    for old in r['lessons']:
        actual=next(l for l in new['lessons'] if l['id']==old['id'])
        if old['slot']//8!=day:assert old['slot']==actual['slot'] and old['parts']==actual['parts']
        elif actual['slot']//8==day:assert all(p['teacher_id']!='t0' for p in actual['parts'])
