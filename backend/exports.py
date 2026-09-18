import csv
import io
import os
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from xml.sax.saxutils import escape
from .domain import slot_stride, bell_label
from .scheduling import indexes, events, DAYS

STATUS={'draft':'Черновик','review':'На проверке','approved':'Утверждено','published':'Опубликовано','archived':'Архив'}
HEADERS=['День','Урок','Время','Классы','Предмет','Тема','Учитель','Кабинет']
def safe(value):
    text=str(value)
    return "'"+text if text.lstrip().startswith(('=','+','-','@','\t','\r','\n')) else text

def rows_for(version, filters):
    data=version['snapshot']; stride=slot_stride(data); ix=indexes(data); ev={e['id']:e for e in events(data,version['week'])}; rows=[]
    for l in sorted(version['lessons'],key=lambda x:(x['slot'],x['id'])):
        e=ev[l['id']]
        if filters.get('class_id') and filters['class_id'] not in e['class_ids']: continue
        if filters.get('subject_id') and filters['subject_id']!=e['subject_id']: continue
        if filters.get('language') and not any(ix['classes'][c]['language']==filters['language'] for c in e['class_ids']): continue
        for pi,p in enumerate(l['parts']):
            if filters.get('teacher_id') and filters['teacher_id']!=p['teacher_id']: continue
            if filters.get('room_id') and filters['room_id']!=p['room_id']: continue
            b=data['settings']['bells'][l['slot']%stride]; course=ix['courses'][e['course_id']]
            topic=course['topics'][e['topic']-1] if e['topic']<=len(course['topics']) else f'Тема {e["topic"]}'
            group=e['parts'][pi]['group_id']
            classes=', '.join(ix['classes'][c]['name'] for c in e['class_ids'])+(f' / {ix["groups"][group]["name"]}' if group else '')
            rows.append([DAYS[l['slot']//stride],bell_label(data,l['slot']%stride),b['start']+'–'+b['end'],classes,ix['subjects'][e['subject_id']]['name'],topic,ix['teachers'][p['teacher_id']]['name'],ix['rooms'][p['room_id']]['name']])
    return rows

def export(version, vid, status, fmt, filters):
    data=version['snapshot']; rows=rows_for(version,filters)
    meta=[data['settings']['school'],version['week'] or 'Базовая неделя',f'Версия {vid}',STATUS[status],datetime.now(timezone.utc).isoformat()]
    output=io.BytesIO()
    if fmt=='csv':
        stream=io.StringIO(newline=''); writer=csv.writer(stream,delimiter=';'); writer.writerow([safe(x) for x in meta]); writer.writerow(HEADERS)
        writer.writerows([[safe(x) for x in row] for row in rows]); return stream.getvalue().encode('utf-8-sig'),'text/csv; charset=utf-8'
    if fmt=='xlsx':
        wb=Workbook(); wb.remove(wb.active)
        def sheet(name,rs):
            ws=wb.create_sheet(name[:31]); ws.append([safe(x) for x in meta]); ws.append(HEADERS)
            for row in rs: ws.append([safe(x) for x in row])
            ws.freeze_panes='D3'; ws.auto_filter.ref=f'A2:H{max(2,ws.max_row)}'
            for cell in ws[2]: cell.font=Font(color='FFFFFF',bold=True); cell.fill=PatternFill('solid',fgColor='24334B')
            for col,width in zip('ABCDEFGH',[20,10,18,20,25,25,30,18]): ws.column_dimensions[col].width=width
            for row in ws:
                for cell in row: cell.alignment=Alignment(vertical='top',wrap_text=True)
        sheet('Общая сетка',rows)
        for kind,key,title in [('classes','class_id','Классы'),('teachers','teacher_id','Учителя'),('rooms','room_id','Кабинеты')]:
            combined=[]
            for r in data[kind]:
                if filters.get(key) and filters[key]!=r['id']: continue
                combined.extend(rows_for(version,{**filters,key:r['id']}))
            sheet(title,combined)
        ws=wb.create_sheet('Конфликты'); ws.append(['Причина','Штраф'])
        for w in version.get('quality',{}).get('warnings',[]): ws.append([safe(w['message']),w['penalty']])
        wb.save(output); return output.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    candidates=[os.getenv('PDF_FONT',''),'C:/Windows/Fonts/arial.ttf','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf']
    font=next((p for p in candidates if p and Path(p).is_file()),None)
    if not font: raise ValueError('Для PDF установите шрифт DejaVu Sans или задайте PDF_FONT')
    if 'School' not in pdfmetrics.getRegisteredFontNames(): pdfmetrics.registerFont(TTFont('School',font))
    styles=getSampleStyleSheet(); style=styles['Normal']; style.fontName='School'; style.fontSize=7; style.leading=10
    doc=SimpleDocTemplate(output,pagesize=landscape(A4),rightMargin=24,leftMargin=24,topMargin=24,bottomMargin=24)
    table=Table([[Paragraph(escape(str(x)),style) for x in row] for row in [HEADERS]+rows],repeatRows=1,colWidths=[75,32,65,65,105,105,150,65])
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e7edf5')),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#cbd5e1')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
    doc.build([Paragraph(escape(' · '.join(meta)),style),Spacer(1,15),table]); return output.getvalue(),'application/pdf'
